"""Tests de l'export Excel du planning hebdomadaire (onglet Créneaux)."""
import os
import sys
import tempfile
import unittest

_test_dir = os.path.dirname(os.path.abspath(__file__))
_src_dir = os.path.join(os.path.dirname(_test_dir), "src") if os.path.basename(_test_dir) == "tests" else _test_dir
if _src_dir not in sys.path:
    sys.path.insert(0, _src_dir)

from infrastructure.planning_excel_export import (
    parse_horaire, clean_group_label, _initiales, encadrants_text,
    assign_lanes, export_planning_excel, export_planning_list_excel,
    fill_for_item,
)


class TestFillForItem(unittest.TestCase):
    def test_autonome_bloc_pastel_orange(self):
        self.assertEqual(fill_for_item({"groupe": "Autonome Bloc", "type": "autonome"}), "FFD8A8")
        self.assertEqual(fill_for_item({"groupe": "autonome  bloc", "type": "autonome"}), "FFD8A8")

    def test_autonome_simple_gris(self):
        self.assertEqual(fill_for_item({"groupe": "Autonome", "type": "autonome"}), "D9D9D9")

    def test_par_type(self):
        self.assertEqual(fill_for_item({"groupe": "Compétition - U11-U13", "type": "compétition"}), "D9EAD3")
        self.assertEqual(fill_for_item({"groupe": "Perfectionnement - U13", "type": "perfectionnement"}), "F8C8DC")


class TestParseHoraire(unittest.TestCase):
    def test_standard_formats(self):
        self.assertEqual(parse_horaire("18:00 - 20:00"), (18 * 60, 20 * 60))
        self.assertEqual(parse_horaire("17:00 18:00"), (17 * 60, 18 * 60))   # sans tiret
        self.assertEqual(parse_horaire("9h30-12h"), (9 * 60 + 30, 12 * 60))

    def test_invalid(self):
        self.assertIsNone(parse_horaire(""))
        self.assertIsNone(parse_horaire("18:00"))
        self.assertIsNone(parse_horaire("20:00 - 18:00"))   # fin avant début


class TestCleanGroupLabel(unittest.TestCase):
    def test_strip_schedule_suffix(self):
        self.assertEqual(clean_group_label("Enfants 2019-2020 (Mercredi 09h30)"), "Enfants 2019-2020")
        self.assertEqual(clean_group_label("Collège (Lundi 18h30)"), "Collège")

    def test_keep_distinctive_suffix(self):
        # Le (1) distingue deux groupes : conservé
        self.assertEqual(clean_group_label("Perfectionnement - U11-U13(1)"), "Perfectionnement - U11-U13(1)")
        self.assertEqual(clean_group_label("Compétition - U11-U13"), "Compétition - U11-U13")


class TestInitiales(unittest.TestCase):
    def test_simple_and_composed(self):
        self.assertEqual(_initiales("Camille DIDIER"), "C.D.")
        self.assertEqual(_initiales("Clément de Lima Ferreira"), "C.D.L.F.")

    def test_empty(self):
        self.assertEqual(_initiales(""), "")
        self.assertEqual(_initiales("   "), "")


class TestEncadrantsText(unittest.TestCase):
    def test_full_names_when_space(self):
        label = "Compétition U11-U13"
        # bloc de 6 lignes : large
        txt = encadrants_text(["Camille DIDIER", "Romain GOETHALS"], 6, label)
        self.assertIn("Camille DIDIER", txt)

    def test_initials_when_narrow(self):
        label = "Compétition U11-U13"
        # bloc de 3 lignes : le libellé prend 2 lignes -> le nom complet ne tient
        # pas, les initiales (1 ligne) oui
        txt = encadrants_text(["Camille DIDIER", "Romain GOETHALS"], 3, label)
        self.assertNotIn("Camille DIDIER", txt)
        self.assertIn("C.D.", txt)
        self.assertIn("R.G.", txt)

    def test_omitted_when_no_room(self):
        label = "Compétition U11-U13"
        # bloc de 2 lignes : libellé sur 2 lignes -> aucun encadrant possible
        txt = encadrants_text(["Camille DIDIER", "Romain GOETHALS"], 2, label)
        self.assertEqual(txt, "")

    def test_empty_list(self):
        self.assertEqual(encadrants_text([], 6, "X"), "")


class TestAssignLanes(unittest.TestCase):
    def test_no_overlap_single_lane(self):
        slots = [{"start": 18 * 60, "end": 20 * 60}, {"start": 20 * 60, "end": 22 * 60}]
        assignments, lanes = assign_lanes(slots)
        self.assertEqual(lanes, 1)
        self.assertEqual([lane for _s, lane in assignments], [0, 0])

    def test_overlap_two_lanes(self):
        slots = [{"start": 18 * 60, "end": 21 * 60}, {"start": 19 * 60, "end": 22 * 60}]
        assignments, lanes = assign_lanes(slots)
        self.assertEqual(lanes, 2)
        self.assertEqual(assignments[0][1], 0)
        self.assertEqual(assignments[1][1], 1)

    def test_same_time_two_lanes(self):
        slots = [{"start": 18 * 60, "end": 22 * 60}, {"start": 18 * 60, "end": 22 * 60}]
        _assignments, lanes = assign_lanes(slots)
        self.assertEqual(lanes, 2)


class TestExportExcel(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self._tmp.name, "planning.xlsx")

    def tearDown(self):
        self._tmp.cleanup()

    def _planning(self):
        return [
            {"groupe": "Compétition - U11-U13", "type": "compétition", "jour": "Lundi",
             "horaires": "18:00 - 20:00", "encadrants": ["Camille DIDIER"]},
            {"groupe": "Autonome", "type": "autonome", "jour": "Lundi",
             "horaires": "20:00 - 22:00", "encadrants": ["Nathalie LEROY"]},
            {"groupe": "Perfectionnement - U11-U13(1)", "type": "perfectionnement", "jour": "Mercredi",
             "horaires": "13:00 - 14:30", "encadrants": ["Thibault Natale"]},
            {"groupe": "Enfants 2019-2020 (Mercredi 09h30)", "type": "cours", "jour": "Mercredi",
             "horaires": "09:30 - 10:30", "encadrants": []},
        ]

    def test_export_creates_grid(self):
        from openpyxl import load_workbook
        out = export_planning_excel(self._planning(), "2026-2027", self.path)
        self.assertTrue(os.path.exists(out))
        wb = load_workbook(out)
        ws = wb.active
        self.assertEqual(ws.title, "Planning")
        # Titre saison
        self.assertEqual(ws.cell(row=1, column=2).value, "Saison 2026 - 2027")
        # En-têtes de jours (1 voie par jour ici : Lundi=col2 … Samedi=col7)
        self.assertEqual(ws.cell(row=2, column=2).value, "LUNDI")
        self.assertEqual(ws.cell(row=2, column=4).value, "MERCREDI")
        self.assertEqual(ws.cell(row=2, column=7).value, "SAMEDI")
        # Libellés horaires gauche/droite
        self.assertEqual(ws.cell(row=3, column=1).value, "9h")
        self.assertEqual(ws.cell(row=3, column=8).value, "9h")
        # Hiérarchie des traits : ligne d'heure = medium, demi-heure = thin
        self.assertEqual(ws.cell(row=3, column=3).border.bottom.style, "thin")    # 9h30
        self.assertEqual(ws.cell(row=4, column=3).border.bottom.style, "medium")  # 10h
        self.assertEqual(ws.cell(row=3, column=3).border.top.style, "medium")     # cadre 9h
        # Séparateurs de jours = thick noir ; voies d'un même jour = thin
        # (1 voie par jour ici : chaque colonne termine un jour)
        for col in (2, 3, 4, 5, 6, 7):
            self.assertEqual(ws.cell(row=3, column=col).border.right.style, "thick", col)
        # Bloc fusionné : Compétition Lundi 18h-20h -> lignes 21..24
        merged = [str(m) for m in ws.merged_cells.ranges]
        self.assertTrue(any(r.startswith("B21") and r.endswith("B24") for r in merged),
                        f"bloc 18h-20h attendu, merged={merged}")
        self.assertTrue(any(r.startswith("B25") and r.endswith("B28") for r in merged),
                        f"bloc 20h-22h attendu, merged={merged}")
        # Contenu du bloc avec encadrant (place suffisante : noms complets)
        self.assertIn("Compétition - U11-U13", ws.cell(row=21, column=2).value)
        self.assertIn("Camille DIDIER", ws.cell(row=21, column=2).value)
        self.assertIn("Autonome", ws.cell(row=25, column=2).value)

    def test_export_without_valid_slots_raises(self):
        with self.assertRaises(ValueError):
            export_planning_excel([{"groupe": "X", "jour": "Lundi", "horaires": "invalide"}], "2026-2027", self.path)

    def test_day_separator_vs_lane_separator(self):
        """Deux créneaux qui se chevauchent : voie thin, fin de jour thick."""
        from openpyxl import load_workbook
        planning = [
            {"groupe": "Compétition - U11-U13", "type": "compétition", "jour": "Lundi",
             "horaires": "18:00 - 20:00", "encadrants": ["Camille DIDIER"]},
            {"groupe": "Collège", "type": "cours", "jour": "Lundi",
             "horaires": "18:30 - 20:00", "encadrants": []},
            {"groupe": "Autonome", "type": "autonome", "jour": "Mardi",
             "horaires": "20:00 - 22:00", "encadrants": []},
        ]
        out = export_planning_excel(planning, "2026-2027", self.path)
        ws = load_workbook(out).active
        # Lundi : 2 voies (col2 = voie 0, col3 = fin du jour) puis Mardi (col4 = fin du jour)
        self.assertEqual(ws.cell(row=3, column=2).border.right.style, "thin")
        self.assertEqual(ws.cell(row=3, column=3).border.right.style, "thick")
        self.assertEqual(ws.cell(row=3, column=4).border.right.style, "thick")


class TestExportListExcel(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self._tmp.name, "planning_liste.xlsx")

    def tearDown(self):
        self._tmp.cleanup()

    def test_list_matches_db_shape(self):
        from openpyxl import load_workbook
        planning = [
            {"id": 7, "groupe": "Compétition - U11-U13", "type": "compétition", "jour": "Mercredi",
             "horaires": "16:00 - 18:00", "encadrants": ["Clément DE LIMA FERREIRA"],
             "helloasso_tarifs": ["Compétition U11-U13"], "categorie_age": "U11-U13",
             "naissance_min": "2014-01-01", "naissance_max": "2016-12-31"},
            {"id": 3, "groupe": "Autonome", "type": "autonome", "jour": "Lundi",
             "horaires": "20:00 - 22:00", "encadrants": ["Nathalie LEROY"],
             "helloasso_tarifs": ["Adultes autonomes"], "categorie_age": "Adultes",
             "naissance_min": "", "naissance_max": ""},
        ]
        out = export_planning_list_excel(planning, "2026-2027", self.path)
        ws = load_workbook(out).active
        self.assertEqual(ws.title, "Planning (liste)")
        # En-têtes proches de la BDD
        headers = [ws.cell(row=1, column=c).value for c in range(1, 10)]
        self.assertEqual(headers[:5], ["Jour", "Groupe", "Type", "Horaires", "Encadrants"])
        # Tri Lundi → Samedi puis heure
        self.assertEqual(ws.cell(row=2, column=1).value, "Lundi")
        self.assertEqual(ws.cell(row=3, column=1).value, "Mercredi")
        self.assertEqual(ws.cell(row=2, column=5).value, "Nathalie LEROY")
        self.assertEqual(ws.cell(row=3, column=6).value, "Compétition U11-U13")
        self.assertEqual(ws.cell(row=3, column=8).value, "2014-01-01")

    def test_empty_planning(self):
        from openpyxl import load_workbook
        out = export_planning_list_excel([], "2026-2027", self.path)
        ws = load_workbook(out).active
        self.assertEqual(ws.max_row, 1)   # en-têtes seuls


if __name__ == "__main__":
    unittest.main()
