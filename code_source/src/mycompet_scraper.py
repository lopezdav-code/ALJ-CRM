import os
import requests
import lxml.html
import pandas as pd
from paths import ROOT_DIR

def scrape_mycompet(progress_callback=None) -> str:
    """
    Scrape les résultats de compétitions FFME depuis https://escalade.online/resultat/
    et génère un fichier Excel consolidé dans exports/MyCompet.xlsx.
    Retourne le chemin d'accès au fichier généré.
    """
    disciplines = ["BLOC", "VITESSE", "DIFFICULTÉ"]
    list_cat = ["U13", "U13", "U15", "U15", "U17", "U17", "U19", "U19", "SÉNIOR", "SÉNIOR", "VÉTÉRAN", "VÉTÉRAN"]
    list_sexe = ["FEMME", "HOMME", "FEMME", "HOMME", "FEMME", "HOMME", "FEMME", "HOMME", "FEMME", "HOMME", "FEMME", "HOMME"]
    
    def detect_criterion(tbl_element, keywords):
        # Récupérer l'élément parent externe
        parent = tbl_element.getparent()
        if parent is not None:
            html_snippet = lxml.html.tostring(parent, encoding="utf-8").decode("utf-8").upper()
        else:
            html_snippet = lxml.html.tostring(tbl_element, encoding="utf-8").decode("utf-8").upper()
            
        if len(html_snippet) < 50:
            html_snippet = lxml.html.tostring(tbl_element, encoding="utf-8").decode("utf-8").upper()
            
        # Nettoyage UTF-8 et accents
        html_snippet = html_snippet.replace("Ã©", "E").replace("É", "E")
        
        for kw in keywords:
            kw_clean = kw.upper().replace("É", "E")
            if kw_clean in html_snippet:
                return kw
        return ""

    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "Accept-Encoding": "identity"
    }
    
    all_rows = []
    header_cols = []
    
    for num_page in range(1, 4):
        disc = disciplines[num_page - 1]
        url = f"https://escalade.online/resultat/classement_{num_page}.html"
        
        if progress_callback:
            progress_callback(f"🌐 Téléchargement de la discipline : {disc}...", 30 + num_page * 15)
            
        response = requests.get(url, headers=headers, timeout=15)
        if response.status_code != 200:
            raise Exception(f"Impossible de charger la page {url} (code {response.status_code})")
            
        # Parser le HTML
        html_text = response.text
        doc = lxml.html.fromstring(html_text)
        tables = doc.xpath("//table")
        
        if progress_callback:
            progress_callback(f"📊 Analyse des données de {disc} ({len(tables)} tableaux)...", 30 + num_page * 15 + 5)
            
        valid_table_count = 0
        for tbl in tables:
            rows = tbl.xpath(".//tr")
            if len(rows) > 2:
                valid_table_count += 1
                
                # 1. Détection de la catégorie et du sexe
                cat = detect_criterion(tbl, ["U13", "U15", "U17", "U19", "SÉNIOR", "SENIOR", "VÉTÉRAN", "VETERAN"])
                sexe = detect_criterion(tbl, ["FEMME", "HOMME", "DAMES", "MESSIEURS"])
                
                # Normalisation
                if cat == "SENIOR": cat = "SÉNIOR"
                if cat == "VETERAN": cat = "VÉTÉRAN"
                if sexe == "DAMES": sexe = "FEMME"
                if sexe == "MESSIEURS": sexe = "HOMME"
                
                # Repli automatique
                if not cat:
                    cat = list_cat[(valid_table_count - 1) % 12]
                if not sexe:
                    sexe = list_sexe[(valid_table_count - 1) % 12]
                    
                # 2. Lecture des lignes du tableau
                for r in rows:
                    # Chercher les td, sinon th
                    cells = r.xpath(".//td")
                    if len(cells) == 0:
                        cells = r.xpath(".//th")
                        is_header_row = True
                    else:
                        is_header_row = False
                        
                    cell_texts = [cell.text_content().strip() for cell in cells]
                    
                    if is_header_row:
                        # Si on n'a pas encore défini les en-têtes globaux de colonnes
                        if not header_cols and len(cell_texts) > 0:
                            header_cols = ["Discipline", "Catégorie", "Sexe"] + cell_texts
                    else:
                        # Ligne de données
                        data_row = [disc, cat, sexe] + cell_texts
                        all_rows.append(data_row)
                        
    if not all_rows:
        raise Exception("Aucune donnée de classement trouvée sur le site.")
        
    # Construire les noms de colonnes si absents
    if not header_cols:
        max_cols = max(len(row) for row in all_rows)
        header_cols = ["Discipline", "Catégorie", "Sexe"] + [f"Col {i}" for i in range(4, max_cols + 1)]
        
    # S'assurer que chaque ligne de données a le même nombre de colonnes que l'en-tête (en complétant par du vide)
    num_headers = len(header_cols)
    for idx, r in enumerate(all_rows):
        if len(r) < num_headers:
            all_rows[idx] = r + [""] * (num_headers - len(r))
        elif len(r) > num_headers:
            all_rows[idx] = r[:num_headers]
            
    # Créer le DataFrame pandas
    df = pd.DataFrame(all_rows, columns=header_cols)
    
    if progress_callback:
        progress_callback("💾 Enregistrement et mise en page du fichier Excel...", 90)
        
    # Enregistrer le résultat
    dest_file = os.path.join(ROOT_DIR, "exports", "MyCompet.xlsx")
    os.makedirs(os.path.dirname(dest_file), exist_ok=True)
    from infrastructure.spreadsheet_safety import safe_to_excel
    safe_to_excel(df, dest_file, index=False)  # scraped data must never become a formula
    
    # Rendre l'ajustement automatique des colonnes avec openpyxl pour faire un joli rendu Excel !
    import openpyxl
    wb = openpyxl.load_workbook(dest_file)
    ws = wb.active
    
    # Rendre l'en-tête gras (Font style)
    from openpyxl.styles import Font
    bold_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold_font
        
    # Auto-fit columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    wb.save(dest_file)
    wb.close()
    
    return dest_file


def normalize_result_url(url_or_id: str) -> str:
    """Transforme un ID numérique ou une URL partielle en URL complète MyCompet."""
    clean = str(url_or_id or "").strip()
    if not clean:
        return ""
    if clean.isdigit():
        return f"https://mycompet.ffme.fr/resultat/resultat_{clean}"
    if clean.startswith("resultat_"):
        return f"https://mycompet.ffme.fr/resultat/{clean}"
    if not clean.startswith("http://") and not clean.startswith("https://"):
        return f"https://mycompet.ffme.fr/resultat/resultat_{clean}"
    return clean


def parse_mycompet_results_html(html_text: str, club_keyword: str = "JONAGE") -> list[dict]:
    """Parse le HTML d'une page de résultats MyCompet et extrait les résultats du club."""
    import re

    if not html_text:
        return []

    doc = lxml.html.fromstring(html_text)
    regroupements = doc.xpath("//div[contains(@class, 'regroupement')]")

    # Si aucun regroupement div trouvé, chercher les tables directement
    elements = regroupements if regroupements else doc.xpath("//table")
    club_kw = club_keyword.upper().strip()

    seen = {}
    for elem in elements:
        # Trouver la catégorie si présent
        cat_name = ""
        header_el = elem.xpath(".//div[contains(@class, 'regroupement-header')]")
        if header_el:
            cat_name = header_el[0].text_content().strip().replace("[", "").replace("]", "")
        if not cat_name and elem.tag == "table":
            # Chercher dans les balises précédentes
            prev = elem.xpath("preceding::div[contains(text(), 'U13') or contains(text(), 'U15') or contains(text(), 'U17') or contains(text(), 'U19') or contains(text(), 'Sénior') or contains(text(), 'Senior') or contains(text(), 'Vétéran') or contains(text(), 'Veteran')]")
            if prev:
                cat_match = re.search(r'(\[?[Uu]\d{2}\]?\s*(?:Femme|Homme|F|H)?|\[?(?:Sénior|Senior|Vétéran|Veteran)\]?\s*(?:Femme|Homme|F|H)?)', prev[-1].text_content())
                cat_name = cat_match.group(1).replace("[", "").replace("]", "").strip() if cat_match else prev[-1].text_content().strip()

        # Calculer le nombre total de participants dans cette catégorie / épreuve
        athlete_rows = elem.xpath(".//table//tr[.//a[contains(@href, 'palmares')]]") if elem.tag != "table" else elem.xpath(".//tr[.//a[contains(@href, 'palmares')]]")
        total_in_cat = len(athlete_rows)

        rows = elem.xpath(".//tr[contains(translate(., 'abcdefghijklmnopqrstuvwxyz', 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'), '" + club_kw + "')]")
        for r in rows:
            a_elem = r.xpath(".//a[contains(@href, 'palmares')]")
            if not a_elem:
                continue
            href = a_elem[0].attrib.get('href', '')
            lic_match = re.search(r'palmares_(\d+)', href)
            licence = lic_match.group(1) if lic_match else ""

            strong = a_elem[0].xpath(".//strong")
            full_name = strong[0].text_content().strip() if strong else a_elem[0].text_content().split("\n")[0].strip()

            parts = full_name.split()
            if len(parts) >= 2:
                nom_parts = [p for p in parts if p.isupper()]
                prenom_parts = [p for p in parts if not p.isupper()]
                if nom_parts and prenom_parts:
                    nom = " ".join(nom_parts)
                    prenom = " ".join(prenom_parts)
                else:
                    nom = parts[0]
                    prenom = " ".join(parts[1:])
            else:
                nom = full_name
                prenom = ""

            cells = [c.text_content().strip() for c in r.xpath(".//td")]
            rank_str = cells[0] if cells else ""
            perf_str = cells[-3] if len(cells) >= 4 else ""

            key = (licence or full_name).lower()
            # Conserver la catégorie la plus précise
            if key not in seen or (cat_name and not seen[key].get("categorie")):
                seen[key] = {
                    "nom": nom,
                    "prenom": prenom,
                    "nom_complet": full_name,
                    "licence": licence,
                    "categorie": cat_name,
                    "classement": rank_str,
                    "total_categorie": total_in_cat,
                    "performance": perf_str,
                }

    return list(seen.values())


def fetch_and_parse_mycompet_results(url_or_id: str, club_keyword: str = "JONAGE") -> list[dict]:
    """Télécharge la page de résultats MyCompet et retourne les résultats extraits pour le club."""
    url = normalize_result_url(url_or_id)
    if not url:
        raise ValueError("URL ou identifiant de résultat MyCompet manquant.")

    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "fr,fr-FR;q=0.9,en;q=0.8"
    }

    resp = requests.get(url, headers=headers, timeout=15)
    if resp.status_code != 200:
        raise RuntimeError(f"Impossible de récupérer les résultats MyCompet depuis {url} (code HTTP {resp.status_code})")

    return parse_mycompet_results_html(resp.text, club_keyword=club_keyword)


def find_mycompet_result_url(comp_nom: str, comp_date: str = "", comp_lieu: str = "", id_ffme: str = "") -> str:
    """Tente de trouver automatiquement l'URL de résultat MyCompet sans saisie manuelle."""
    import unicodedata
    import re

    def _clean(t: str) -> str:
        t = unicodedata.normalize('NFKD', t or "")
        t = "".join(c for c in t if not unicodedata.combining(c))
        return re.sub(r'[^a-zA-Z0-9]+', ' ', t).lower().strip()

    clean_nom = _clean(comp_nom)
    clean_lieu = _clean(comp_lieu)

    stop_words = {'coupe', 'regional', 'regionale', 'departemental', 'departementale', 'bloc', 'difficulte', 'etape', 'open', 'france', 'championnat'}

    # 1. Vérification dans la page d'archives officielles FFME
    try:
        url_archives = "https://www.ffme.fr/escalade/competition/archives-resultats-escalade/"
        resp = requests.get(url_archives, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}, timeout=10)
        if resp.status_code == 200:
            doc = lxml.html.fromstring(resp.text)
            links = doc.xpath("//a[contains(@href, 'mycompet.ffme.fr/resultat/resultat_')]")

            best_match = None
            for a in links:
                href = a.attrib.get('href', '')
                parent = a.getparent()
                full_text = " ".join(parent.text_content().split())
                clean_full = _clean(full_text)

                words = [w for w in clean_nom.split() if len(w) > 3 and w not in stop_words]
                if words:
                    matches = sum(1 for w in words if w in clean_full)
                    if matches >= max(1, len(words) - 1):
                        return href

                if clean_nom and (clean_nom in clean_full or clean_full in clean_nom):
                    return href

                if clean_lieu and clean_lieu in clean_full:
                    best_match = href

            if best_match:
                return best_match
    except Exception as e:
        print(f"⚠️ [MYCOMPET] Erreur recherche automatique archives FFME : {e}")

    # 2. Si non trouvé dans les archives, sonder les pages MyCompet récentes
    try:
        for offset in range(13710, 13640, -1):
            url_test = f"https://mycompet.ffme.fr/resultat/resultat_{offset}"
            try:
                r_probe = requests.get(url_test, headers={"User-Agent": "Mozilla/5.0"}, timeout=3)
                if r_probe.status_code == 200:
                    probe_doc = lxml.html.fromstring(r_probe.text)
                    divs = probe_doc.xpath("//div[contains(@class, 'cr-Titre') or contains(@class, 'title')]")
                    probe_title = _clean(divs[0].text_content() if divs else "")
                    if clean_nom and (clean_nom in probe_title or probe_title in clean_nom):
                        return url_test
                    words = [w for w in clean_nom.split() if len(w) > 3 and w not in stop_words]
                    if words and sum(1 for w in words if w in probe_title) >= max(1, len(words) - 1):
                        return url_test
            except Exception:
                continue
    except Exception as e:
        print(f"⚠️ [MYCOMPET] Erreur sonde récente MyCompet : {e}")

    return ""
